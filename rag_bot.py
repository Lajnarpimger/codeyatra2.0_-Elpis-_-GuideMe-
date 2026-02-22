import os
import re
import base64
import shutil
from pathlib import Path
from typing import List, Dict, Any, Tuple

import streamlit as st
import chromadb
from chromadb.utils import embedding_functions
from pypdf import PdfReader
from pptx import Presentation
import ollama

# ==============================
# Config
# ==============================
DB_DIR = "chroma_db"
COLLECTION_NAME = "kbase"  # must be 3+ chars for some chroma versions
EMBED_MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"

TEXT_MODEL = "llama3.2"
VISION_MODEL = "llava"

CHUNK_SIZE = 1100
CHUNK_OVERLAP = 180
TOP_K = 6

UPLOAD_DIR = Path("kb_uploads")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

SUPPORTED_EXTS = {".pdf", ".pptx", ".txt", ".md", ".png", ".jpg", ".jpeg", ".webp"}


# ==============================
# Text extraction
# ==============================
def read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")

def read_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"[PDF_PAGE {i+1}]\n{text}")
    return "\n\n".join(parts)

def read_pptx_text(path: Path) -> str:
    prs = Presentation(str(path))
    parts = []
    for si, slide in enumerate(prs.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                t = shape.text.strip()
                if t:
                    slide_text.append(t)
        if slide_text:
            parts.append(f"[PPTX_SLIDE {si}]\n" + "\n".join(slide_text))
    return "\n\n".join(parts)

def llava_image_to_text(image_path: Path) -> str:
    b64 = base64.b64encode(image_path.read_bytes()).decode("utf-8")
    prompt = (
        "Extract all readable text from this image. "
        "If it's a slide/syllabus, return clean structured text. "
        "Do not add extra commentary."
    )
    resp = ollama.chat(
        model=VISION_MODEL,
        messages=[{"role": "user", "content": prompt, "images": [b64]}],
    )
    return (resp["message"]["content"] or "").strip()

def detect_type(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in [".txt", ".md"]:
        return "text"
    if ext == ".pdf":
        return "pdf"
    if ext == ".pptx":
        return "pptx"
    if ext in [".png", ".jpg", ".jpeg", ".webp"]:
        return "image"
    return "unknown"

def clean_text(s: str) -> str:
    s = s.replace("\r", "\n")
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()

def chunk_text(text: str, chunk_size: int, overlap: int) -> List[str]:
    if not text:
        return []
    chunks = []
    start = 0
    n = len(text)
    while start < n:
        end = min(n, start + chunk_size)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == n:
            break
        start = max(0, end - overlap)
    return chunks


# ==============================
# Vector DB
# ==============================
def get_collection():
    client = chromadb.PersistentClient(path=DB_DIR)
    emb_fn = embedding_functions.SentenceTransformerEmbeddingFunction(
        model_name=EMBED_MODEL_NAME
    )
    return client.get_or_create_collection(
        name=COLLECTION_NAME,
        embedding_function=emb_fn,
        metadata={"hnsw:space": "cosine"},
    )

def reset_db():
    # deletes persistent db folder
    if Path(DB_DIR).exists():
        shutil.rmtree(DB_DIR, ignore_errors=True)


# ==============================
# Ingest
# ==============================
def ingest_files(file_paths: List[Path]) -> Dict[str, Any]:
    col = get_collection()
    docs, metas, ids = [], [], []
    report = {"ingested": [], "skipped": [], "errors": []}

    for fp in file_paths:
        try:
            ftype = detect_type(fp)
            if ftype == "unknown":
                report["skipped"].append(str(fp))
                continue

            if ftype == "text":
                raw = read_txt(fp)
            elif ftype == "pdf":
                raw = read_pdf_text(fp)
            elif ftype == "pptx":
                raw = read_pptx_text(fp)
            elif ftype == "image":
                raw = llava_image_to_text(fp)
            else:
                report["skipped"].append(str(fp))
                continue

            raw = clean_text(raw)
            chunks = chunk_text(raw, CHUNK_SIZE, CHUNK_OVERLAP)

            for i, ch in enumerate(chunks):
                cid = f"{fp.name}::chunk{i}"
                ids.append(cid)
                docs.append(ch)
                metas.append({
                    "source_file": str(fp),
                    "file_name": fp.name,
                    "chunk_index": i,
                    "file_type": ftype,
                })

            report["ingested"].append({"file": fp.name, "chunks": len(chunks), "type": ftype})

        except Exception as e:
            report["errors"].append({"file": str(fp), "error": str(e)})

    if docs:
        col.upsert(documents=docs, metadatas=metas, ids=ids)

    report["total_chunks"] = len(docs)
    return report


# ==============================
# Retrieval + Answer
# ==============================
def retrieve(query: str, k: int = TOP_K) -> List[Dict[str, Any]]:
    col = get_collection()
    res = col.query(query_texts=[query], n_results=k)
    out = []
    for doc, meta, _id in zip(res["documents"][0], res["metadatas"][0], res["ids"][0]):
        out.append({"id": _id, "text": doc, "meta": meta})
    return out

def answer_from_kb(question: str) -> Tuple[str, List[Dict[str, Any]]]:
    hits = retrieve(question, TOP_K)
    if not hits:
        return "I don’t know based on the provided files.", []

    blocks = []
    for h in hits:
        blocks.append(f"[{h['id']} | {h['meta']['file_name']}]\n{h['text']}")
    context = "\n\n---\n\n".join(blocks)

    system = (
        "You are a RAG assistant. Answer ONLY using the provided context.\n"
        "Rules:\n"
        "1) If the context does not contain the answer, say: "
        "\"I don’t know based on the provided files.\".\n"
        "2) Do NOT use outside knowledge.\n"
        "3) Provide citations as chunk IDs like [file::chunkN].\n"
        "4) Be concise and correct.\n"
    )

    user = f"CONTEXT:\n{context}\n\nQUESTION:\n{question}\n\nAnswer with citations."

    resp = ollama.chat(
        model=TEXT_MODEL,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
    )
    return (resp["message"]["content"] or "").strip(), hits


# ==============================
# Streamlit UI
# ==============================
st.set_page_config(page_title="Local RAG Bot", layout="wide")
st.title("📚 Local RAG Bot (PDF/PPTX/TXT/Image → Answer)")

with st.sidebar:
    st.header("Settings")
    st.write("Models (Ollama):")
    st.code(f"Text: {TEXT_MODEL}\nVision: {VISION_MODEL}\nEmbeddings: {EMBED_MODEL_NAME}")
    st.write("DB:", DB_DIR, "| Collection:", COLLECTION_NAME)

    if st.button("🧹 Reset knowledge base (delete DB)"):
        reset_db()
        st.success("Deleted DB. Re-ingest your files.")

st.subheader("1) Upload files")
uploaded = st.file_uploader(
    "Upload PDF, PPTX, TXT/MD, or images (PNG/JPG/WebP).",
    type=[e[1:] for e in SUPPORTED_EXTS],
    accept_multiple_files=True
)

col1, col2 = st.columns([1, 1])
with col1:
    ingest_clicked = st.button("📥 Ingest uploaded files")
with col2:
    show_files = st.checkbox("Show uploaded file list", value=True)

saved_paths: List[Path] = []
if uploaded:
    for uf in uploaded:
        save_path = UPLOAD_DIR / uf.name
        with open(save_path, "wb") as f:
            f.write(uf.getbuffer())
        saved_paths.append(save_path)

if show_files:
    st.write("Files in upload folder:")
    existing = sorted([p.name for p in UPLOAD_DIR.glob("*") if p.is_file()])
    if existing:
        st.code("\n".join(existing))
    else:
        st.info("No files uploaded yet.")

if ingest_clicked:
    if not saved_paths:
        # also ingest anything already in UPLOAD_DIR
        saved_paths = [p for p in UPLOAD_DIR.glob("*") if p.is_file()]

    if not saved_paths:
        st.warning("Upload at least one file first.")
    else:
        with st.spinner("Ingesting… (PDF/PPTX text extraction + images via LLaVA)"):
            rep = ingest_files(saved_paths)
        st.success(f"Done. Total chunks added/updated: {rep['total_chunks']}")
        st.json(rep)

st.divider()
st.subheader("2) Ask a question")

if "chat" not in st.session_state:
    st.session_state.chat = []

question = st.text_input("Your question", placeholder="e.g., Summarize the syllabus grading policy.")
ask = st.button("💬 Ask")

if ask and question.strip():
    with st.spinner("Thinking (RAG)…"):
        ans, hits = answer_from_kb(question.strip())

    st.session_state.chat.append({"q": question.strip(), "a": ans, "hits": hits})

for item in reversed(st.session_state.chat):
    st.markdown(f"**You:** {item['q']}")
    st.markdown(f"**Bot:** {item['a']}")
    if item["hits"]:
        with st.expander("🔎 Retrieved context (citations)"):
            for h in item["hits"]:
                st.markdown(f"**{h['id']}** — *{h['meta']['file_name']}*")
                st.code(h["text"][:1500] + ("…" if len(h["text"]) > 1500 else ""))
    st.divider()

st.caption("Tip: If it says “I don’t know…”, your answer probably isn’t in the uploaded files or retrieval TOP_K is too low.")